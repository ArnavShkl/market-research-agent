import pandas as pd, numpy as np, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment
D='JioMausam_Germany_Data/csv/'
wb=Workbook()
F=lambda **k: Font(name='Arial',size=k.pop('size',10),**k)
BLUE=F(color='0000FF'); GREEN=F(color='008000'); BOLD=F(bold=True); H1=F(bold=True,size=14); H2=F(bold=True,size=11); MUTED=F(color='666666',italic=True)
YEL=PatternFill('solid',fgColor='FFFF00'); HEAD=PatternFill('solid',fgColor='DCE6F1'); KEY=PatternFill('solid',fgColor='FFF2CC')
thin=Side(style='thin',color='BFBFBF')
def head(ws,row,cols,start=1):
    for j,t in enumerate(cols):
        c=ws.cell(row,start+j,t); c.font=BOLD; c.fill=HEAD; c.alignment=Alignment(wrap_text=True,vertical='top')
def widths(ws,w):
    for i,x in enumerate(w,1): ws.column_dimensions[L(i)].width=x
def setfont(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.font is None or c.font.name!='Arial': c.font=F(bold=c.font.b if c.font else False,color=c.font.color if c.font and c.font.color else None)
INR='[$₹-4009]#,##0'; INR2='[$₹-4009]#,##0.00'; CR='[$₹-4009]#,##0.0" Cr"'

# ---------- Read me ----------
ws=wb.active; ws.title='Read me'
lines=[('JioMausam: every number, recalculated from the data',H1),
 ('This workbook recomputes the numbers in our pitch with live formulas. Change a yellow cell and everything that depends on it updates.',None),
 ('',None),('Colour code',H2),('Blue text: an input you can change (yellow fill marks the key ones).',BLUE),('Black text: a formula.',None),('Green text: a value pulled from another sheet.',GREEN),
 ('',None),('Sheets',H2),
 ('Calibration: the one scale factor (K = 1.17), fitted on the tuning days 10–12 May 2018 only.',None),
 ('Network hourly: average rain over all 500 links, ours vs radar, for every hour. The correlation on the unseen test days is the 0.95.',None),
 ('Link scores: each link scored on its own on the unseen days. The median is the 0.88; the share above 0.7 is the 91%.',None),
 ('Payout test: 14,256 decisions (1,782 squares × 8 test days). Change the trigger and watch events caught and false triggers move.',None),
 ('India daily: 27,103 days of IMD station reports for Mumbai (Santacruz), Nagpur and Ahmedabad, 2000–2024.',None),
 ('Prices: trigger days counted from the India data, then price = days paid × payout ÷ payout share ÷ 12.',None),
 ('Stress test: every year from 2000 to 2024 replayed as if we had sold the product that year.',None),
 ('Business case: premiums, payouts and Jio’s fees for Kavach Plus and Kavach Basic.',None),
 ('',None),('Sources',H2),
 ('Microwave links and radar: pycomlink example dataset (github.com/pycomlink/pycomlink, folder pycomlink/io/example_data; BSD licence). 500 commercial microwave links in Germany, 1-minute signal levels, 10–20 May 2018. Reference: German Weather Service RADOLAN-YW radar. Link coordinates in the file are shifted for anonymity.',None),
 ('Rain values in Network hourly, Calibration, Link scores and Payout test were computed by prototype/pipeline/rain_pipeline.py from those files. The formulas here redo the scoring.',None),
 ('India: IMD synoptic station reports via NOAA Global Summary of the Day (noaa-gsod-pds.s3.amazonaws.com). Units converted (inches to mm, °F to °C); years with fewer than 330 reports or fewer than 110 monsoon-season reports dropped.',None)]
for i,(t,f) in enumerate(lines,1):
    c=ws.cell(i,1,t); c.font=f or F(); c.alignment=Alignment(wrap_text=True,vertical='top')
ws.column_dimensions['A'].width=120

# ---------- Calibration ----------
lk=pd.read_csv(D+'links.csv')
ws=wb.create_sheet('Calibration')
ws['A1']='Calibration: one scale factor for the whole network, from the tuning days only'; ws['A1'].font=H1
ws['A2']='Our raw estimate tends to read low. K scales it so that, over 10–12 May, our total rain equals radar’s total. The test days (13–20 May) are never used here.'; ws['A2'].font=MUTED
ws['A4']='Radar total, tuning days (mm·h summed over all links)'; ws['B4']='=SUM(C9:C508)'
ws['A5']='Our uncalibrated total, tuning days'; ws['B5']='=SUM(B9:B508)'
ws['A6']='K = radar total ÷ our total'; ws['B6']='=B4/B5'; ws['B6'].number_format='0.000'; ws['A6'].font=BOLD; ws['B6'].font=BOLD; ws['B6'].fill=KEY
for r in (4,5): ws[f'B{r}'].number_format='#,##0.0'
head(ws,8,['Link','Ours, uncalibrated: sum of hourly rain on tuning days (mm)','Radar along the link: same sum (mm)'])
for i,r in lk.iterrows():
    ws.cell(9+i,1,int(r.link_id)); ws.cell(9+i,2,float(r.ours_uncal_mm_tuning)); ws.cell(9+i,3,float(r.radar_mm_tuning))
    ws.cell(9+i,2).number_format='0.00'; ws.cell(9+i,3).number_format='0.00'
widths(ws,[52,34,30]); ws.freeze_panes='A9'

# ---------- Network hourly ----------
net=pd.read_csv(D+'network_average_hourly.csv')
ws=wb.create_sheet('Network hourly')
ws['A1']='Network average, hour by hour: our estimate vs radar'; ws['A1'].font=H1
ws['A2']='Each row averages all 500 links for one hour. Rows 6–77 are tuning hours (10–12 May); rows 78–269 are the unseen test hours (13–20 May).'; ws['A2'].font=MUTED
ws['F4']='Scale factor K'; ws['G4']="=Calibration!B6"; ws['G4'].font=GREEN; ws['G4'].number_format='0.000'
ws['F5']='Correlation, test hours (the 0.95)'; ws['G5']='=CORREL(D78:D269,E78:E269)'; ws['G5'].number_format='0.000'; ws['F5'].font=BOLD; ws['G5'].font=BOLD; ws['G5'].fill=KEY
ws['F6']='Correlation, tuning hours'; ws['G6']='=CORREL(D6:D77,E6:E77)'; ws['G6'].number_format='0.000'
ws['F7']='Our total ÷ radar total, test hours'; ws['G7']='=SUM(D78:D269)/SUM(E78:E269)'; ws['G7'].number_format='0.00'
ws['F9']='Why correlation ignores K: multiplying every value by the same number changes the size, not the up-and-down pattern. K fixes the totals; the correlation was already there.'; ws['F9'].font=MUTED; ws['F9'].alignment=Alignment(wrap_text=True,vertical='top')
ws.merge_cells('F9:H13')
head(ws,5,['Hour (UTC)','Period','Ours, uncalibrated (mm/h)','Ours, calibrated (mm/h)','Radar (mm/h)'])
for i,r in net.iterrows():
    R=6+i; ws.cell(R,1,r.hour_utc); ws.cell(R,2,r.period); ws.cell(R,3,float(r.ours_uncalibrated)); ws.cell(R,4,f'=C{R}*$G$4'); ws.cell(R,5,float(r.radar))
    for cc in (3,4,5): ws.cell(R,cc).number_format='0.000'
widths(ws,[18,10,16,16,14,36,12,12]); ws.freeze_panes='A6'

# ---------- Link scores ----------
ws=wb.create_sheet('Link scores')
ws['A1']='Each link on its own, scored on the unseen days'; ws['A1'].font=H1
ws['A2']='Correlation between our hourly rain and radar along the same path, 13–20 May, computed by rain_pipeline.py. Links with too little rain on those days are not scored.'; ws['A2'].font=MUTED
ws['J4']='Links scored'; ws['K4']='=COUNT(E9:E508)'
ws['J5']='Median correlation (the 0.88)'; ws['K5']='=MEDIAN(E9:E508)'; ws['K5'].number_format='0.000'; ws['K5'].font=BOLD; ws['K5'].fill=KEY; ws['J5'].font=BOLD
ws['J6']='Share of links above'; ws['K6']=0.7; ws['K6'].font=BLUE; ws['K6'].fill=YEL
ws['J7']='Share (the 91%)'; ws['K7']='=COUNTIF(E9:E508,">"&K6)/K4'; ws['K7'].number_format='0.0%'
head(ws,8,['Link','Length (km)','Frequency (GHz)','Polarization','Correlation on unseen days','Our rain, test days (mm)','Radar, test days (mm)','Wet threshold, channel 1 (dB)'])
for i,r in lk.iterrows():
    R=9+i; vals=[int(r.link_id),float(r.length_km),float(r.freq_ch1_GHz),r.polarization_ch1,(None if np.isnan(r.r_unseen_days) else float(r.r_unseen_days)),float(r.ours_mm_unseen),float(r.radar_mm_unseen),float(r.wet_threshold_ch1_dB)]
    for j,v in enumerate(vals,1):
        c=ws.cell(R,j,v)
    for j,fm in [(2,'0.0'),(3,'0.0'),(5,'0.000'),(6,'0.0'),(7,'0.0'),(8,'0.000')]: ws.cell(R,j).number_format=fm
widths(ws,[8,11,13,12,16,16,16,16,4,28,10]); ws.freeze_panes='A9'

# ---------- Payout test ----------
pdx=pd.read_csv(D+'payout_decisions_test_days.csv')
ws=wb.create_sheet('Payout test')
ws['A1']='Payout test: our pay / no-pay decision vs radar’s, per square per day'; ws['A1'].font=H1
ws['A2']='A square pays on a day if any 3-hour window that day had at least the trigger amount of rain. Columns C and D hold the wettest 3-hour window from our map and from radar.'; ws['A2'].font=MUTED
ws['A4']='Trigger: rain in any 3 hours (mm)'; ws['B4']=10; ws['B4'].font=BLUE; ws['B4'].fill=YEL
lab=[('Decisions (square-days)','=COUNT(C16:C14271)',None),('Both pay (hit)','=COUNTIFS(E16:E14271,1,F16:F14271,1)',None),('Radar pays, we miss','=COUNTIFS(E16:E14271,0,F16:F14271,1)',None),
     ('We pay, radar doesn’t (false trigger)','=COUNTIFS(E16:E14271,1,F16:F14271,0)',None),('Neither pays','=COUNTIFS(E16:E14271,0,F16:F14271,0)',None),
     ('Agreement (the 94.6%)','=(B6+B9)/B5','0.0%'),('Radar events we caught (the 68%)','=B6/(B6+B7)','0.0%'),('Our triggers radar didn’t confirm (the 18%)','=B8/(B6+B8)','0.0%')]
for k,(t,f,fm) in enumerate(lab):
    ws.cell(5+k,1,t); c=ws.cell(5+k,2,f)
    if fm: c.number_format=fm; c.font=BOLD; c.fill=KEY
head(ws,15,['Square','Day','Ours: wettest 3 h (mm)','Radar: wettest 3 h (mm)','We pay','Radar pays','Outcome'])
for i,r in enumerate(pdx.itertuples()):
    R=16+i
    ws.cell(R,1,int(r.square_id)); ws.cell(R,2,r.day); ws.cell(R,3,float(r.ours_max_3h_rain_mm)); ws.cell(R,4,float(r.radar_max_3h_rain_mm))
    ws.cell(R,5,f'=IF(C{R}>=$B$4,1,0)'); ws.cell(R,6,f'=IF(D{R}>=$B$4,1,0)')
    ws.cell(R,7,f'=IF(AND(E{R}=1,F{R}=1),"Both pay",IF(F{R}=1,"Missed",IF(E{R}=1,"False trigger","Neither")))')
    ws.cell(R,3).number_format='0.00'; ws.cell(R,4).number_format='0.00'
widths(ws,[40,12,14,14,9,10,14]); ws.freeze_panes='A16'

# ---------- India daily ----------
ind=pd.read_csv('india_daily_imd.csv')
ws=wb.create_sheet('India daily')
head(ws,1,['City','Station','Date','Year','Rain (mm)','Highest temperature (°C)'])
for i,r in enumerate(ind.itertuples()):
    R=2+i; ws.cell(R,1,r.city); ws.cell(R,2,str(r.station)); ws.cell(R,3,r.date); ws.cell(R,4,int(r.year))
    if not np.isnan(r.rain_mm): ws.cell(R,5,float(r.rain_mm)); ws.cell(R,5).number_format='0.0'
    if not np.isnan(r.max_temp_C): ws.cell(R,6,float(r.max_temp_C)); ws.cell(R,6).number_format='0.0'
NI=1+len(ind)
widths(ws,[12,14,12,8,10,14]); ws.freeze_panes='A2'
RNG=lambda col: f"'India daily'!${col}$2:${col}${NI}"

# ---------- Prices ----------
ws=wb.create_sheet('Prices')
ws['A1']='Prices from 25 years of IMD records'; ws['A1'].font=H1
ws['A2']='Count the trigger days in each year, cap them, average them, then price = average days paid × payout ÷ payout share ÷ 12.'; ws['A2'].font=MUTED
ws['A4']='Payout per trigger day (₹)'; ws['B4']=300
ws['A5']='Target payout share (claims ÷ premiums)'; ws['B5']=0.6; ws['B5'].number_format='0%'
ws['A6']='Price floor for Kavach Basic (₹ a month)'; ws['B6']=29
for r in (4,5,6): ws[f'B{r}'].font=BLUE; ws[f'B{r}'].fill=YEL
ws['B4'].number_format=INR; ws['B6'].number_format=INR
ws['C5']='The other 40% pays the insurer, reinsurance, Jio’s fees and running costs.'; ws['C5'].font=MUTED
ws['C6']='₹29 comes from our design rule: a typical location triggers about 0.7 days a year (0.7 × ₹300 ÷ 0.6 ÷ 12 = ₹29).'; ws['C6'].font=MUTED
P=[('Kavach Plus · Mumbai riders','Mumbai','Rain',64.5,5,'No'),('Kavach Plus · 3-day cap','Mumbai','Rain',64.5,3,'No'),('Kavach Basic · Mumbai','Mumbai','Rain',204.5,5,'Yes'),
   ('Kavach Basic · Nagpur heat','Nagpur','Heat',47,5,'Yes'),('Kavach Basic · Ahmedabad heat','Ahmedabad','Heat',47,5,'Yes')]
head(ws,8,['Product','City','Weather (Rain or Heat)','Trigger (mm a day, or °C)','Most days paid a year','Price floor applies?','Avg trigger days a year','Avg days paid a year','Claims per person a year (₹)','Cost price per month (₹)','Price charged per month (₹)'])
YR0=16  # year table starts
for k,p in enumerate(P):
    R=9+k
    for j,v in enumerate(p,1):
        c=ws.cell(R,j,v)
        if j>=2: c.font=BLUE
    dc=L(2+2*k); pc=L(3+2*k)  # per-year columns
    ws.cell(R,7,f'=AVERAGE({dc}{YR0+1}:{dc}{YR0+25})'); ws.cell(R,8,f'=AVERAGE({pc}{YR0+1}:{pc}{YR0+25})')
    ws.cell(R,9,f'=H{R}*$B$4'); ws.cell(R,10,f'=I{R}/$B$5/12'); ws.cell(R,11,f'=IF(F{R}="Yes",MAX(J{R},$B$6),J{R})')
    ws.cell(R,7).number_format='0.00'; ws.cell(R,8).number_format='0.00'; ws.cell(R,9).number_format=INR; ws.cell(R,10).number_format=INR2; ws.cell(R,11).number_format=INR
    ws.cell(R,11).font=BOLD; ws.cell(R,11).fill=KEY
ws.cell(YR0-1,1,'Trigger days and days paid, year by year').font=H2
ws.cell(YR0,1,'Year').font=BOLD; ws.cell(YR0,1).fill=HEAD
for k,p in enumerate(P):
    a=ws.cell(YR0,2+2*k,f'{p[0]}: trigger days'); b=ws.cell(YR0,3+2*k,'days paid')
    for c in (a,b): c.font=BOLD; c.fill=HEAD; c.alignment=Alignment(wrap_text=True,vertical='top')
for y in range(25):
    R=YR0+1+y; ws.cell(R,1,str(2000+y))
    for k in range(len(P)):
        pr=9+k; dc=2+2*k
        ws.cell(R,dc,f'=IF($C${pr}="Rain",COUNTIFS({RNG("A")},$B${pr},{RNG("D")},VALUE($A{R}),{RNG("E")},">="&$D${pr}),COUNTIFS({RNG("A")},$B${pr},{RNG("D")},VALUE($A{R}),{RNG("F")},">="&$D${pr}))')
        ws.cell(R,dc+1,f'=MIN({L(dc)}{R},$E${pr})')
widths(ws,[34,14,14,14,12,12,12,12,14,14,14]); ws.freeze_panes='A9'

# ---------- Stress test ----------
ws=wb.create_sheet('Stress test')
ws['A1']='Stress test: every year replayed'; ws['A1'].font=H1
ws['A2']='Claims ÷ premiums for each year, if we had sold the product at its cost price. Above 100% the year loses money and reinsurance pays the difference.'; ws['A2'].font=MUTED
head(ws,4,['Product','Premium per person a year (₹)','Worst claims ÷ premiums','Years above 100%','Most trigger days in one year','Year it happened'])
for k,p in enumerate(P):
    R=5+k; col=L(2+k)
    ws.cell(R,1,f'=Prices!A{9+k}').font=GREEN
    ws.cell(R,2,f'=Prices!J{9+k}*12').font=GREEN; ws.cell(R,2).number_format=INR
    ws.cell(R,3,f'=MAX({col}14:{col}38)'); ws.cell(R,3).number_format='0%'; ws.cell(R,3).fill=KEY; ws.cell(R,3).font=BOLD
    ws.cell(R,4,f'=COUNTIF({col}14:{col}38,">1")')
    dcol=L(2+2*k)
    ws.cell(R,5,f'=MAX(Prices!{dcol}17:{dcol}41)'); ws.cell(R,6,f'=INDEX(Prices!A17:A41,MATCH(E{R},Prices!{dcol}17:{dcol}41,0))')
ws.cell(12,1,'Claims ÷ premiums, year by year').font=H2
ws.cell(13,1,'Year').font=BOLD; ws.cell(13,1).fill=HEAD
for k,p in enumerate(P):
    c=ws.cell(13,2+k,p[0]); c.font=BOLD; c.fill=HEAD; c.alignment=Alignment(wrap_text=True,vertical='top')
for y in range(25):
    R=14+y; ws.cell(R,1,str(2000+y))
    for k in range(len(P)):
        pc=L(3+2*k); c=ws.cell(R,2+k,f'=IF($B${5+k}>0,Prices!{pc}{17+y}*Prices!$B$4/$B${5+k},0)'); c.number_format='0%'
widths(ws,[32,18,18,18,18,14]); ws.freeze_panes='A5'

# ---------- Business case ----------
ws=wb.create_sheet('Business case')
ws['A1']='Business case: premiums are not Jio’s revenue'; ws['A1'].font=H1
ws['A2']='A licensed insurer holds the premiums and carries the risk. Jio earns distribution and data fees.'; ws['A2'].font=MUTED
rows=[('Inputs',None,None,H2),
 ('Payout per trigger day (₹)','=Prices!B4',INR,GREEN),('Payout share (claims ÷ premiums)','=Prices!B5','0%',GREEN),
 ('Jio’s fee share, low','0.15','0%',BLUE),('Jio’s fee share, high','0.25','0%',BLUE),
 ('',None,None,None),('Kavach Plus · one delivery app',None,None,H2),
 ('Riders covered','50000','#,##0',BLUE),('Price per rider per month (₹)','=Prices!K9',INR,GREEN),
 ('Premiums a year (₹ crore)','=B10*B11*12/10^7','0.00',None),('Back to riders as claims (₹ crore)','=B12*B5','0.00',None),
 ('Jio’s fees, low (₹ crore)','=B12*B6','0.00',None),('Jio’s fees, high (₹ crore)','=B12*B7','0.00',None),
 ('One city-wide flood day: every rider paid once (₹ crore)','=B10*B4/10^7','0.00',None),('Most one rider can get in a year (₹)','=Prices!E9*B4',INR,None),
 ('',None,None,None),('Kavach Basic · at recharge',None,None,H2),
 ('Jio users (crore)','50','0',BLUE),('Share who buy','0.01','0.0%',BLUE),('Price per person per month (₹)','=Prices!B6',INR,GREEN),
 ('Policies','=B20*10^7*B21','#,##0',None),('Premiums a year (₹ crore)','=B23*B22*12/10^7','0.0',None),('Back to people as claims (₹ crore)','=B24*B5','0.0',None),
 ('Jio’s fees, low (₹ crore)','=B24*B6','0.0',None),('Jio’s fees, high (₹ crore)','=B24*B7','0.0',None),
 ('',None,None,None),('Ahmedabad example: 1 lakh customers',None,None,H2),
 ('Customers','100000','#,##0',BLUE),('Premiums a year (₹ crore)','=B30*Prices!K13*12/10^7','0.00',None),('Expected claims a year (₹ crore)','=B30*Prices!H13*B4/10^7','0.00',None)]
for i,(t,v,fm,f) in enumerate(rows,3):
    ws.cell(i,1,t)
    if f is H2: ws.cell(i,1).font=H2; continue
    if v is None: continue
    c=ws.cell(i,2,float(v) if v.replace('.','',1).isdigit() else v)
    if fm: c.number_format=fm
    if f is not None: c.font=f
    if f is BLUE: c.fill=YEL
ws['C20']='Approximate, as used in our pitch.'; ws['C20'].font=MUTED
ws['C21']='Illustrative adoption, not a forecast.'; ws['C21'].font=MUTED
ws['C22']='The Basic price floor (₹29). City prices run up to ₹52.'; ws['C22'].font=MUTED
for r in (12,24): ws[f'B{r}'].font=BOLD; ws[f'B{r}'].fill=KEY
widths(ws,[52,16,52])
for s in wb.worksheets:
    for row in s.iter_rows():
        for c in row:
            if c.value is not None and (c.font is None or c.font.name!='Arial'):
                c.font=Font(name='Arial',size=c.font.sz or 10,bold=c.font.b,italic=c.font.i,color=c.font.color)
wb.save('JioMausam_Numbers_Workbook.xlsx'); print('saved')
